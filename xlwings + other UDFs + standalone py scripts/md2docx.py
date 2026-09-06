from docx import Document
from docx.shared import Pt
import pathlib

doc = Document()
doc.styles['Normal'].font.name = 'Consolas'
doc.styles['Normal'].font.size = Pt(10)

# FULL PATH
md_path = r'D:\Downloads\WhatsApp_Chat_with_Meta_AI.md'
out_path = r'D:\Downloads\WhatsApp_Chat.docx'

md = pathlib.Path(md_path).read_text(encoding='utf-8', errors='ignore')

for line in md.splitlines():
    if 'You' in line and line.startswith('###'):
        doc.add_heading('You', 3)
    elif 'Meta AI' in line and line.startswith('###'):
        doc.add_heading('Meta AI', 3)
    elif line.strip() == '---':
        doc.add_paragraph('─'*60)
    else:
        doc.add_paragraph(line)

doc.save(out_path)
print(f'Saved {out_path}')
